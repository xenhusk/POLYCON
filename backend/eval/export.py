"""
Export evaluation results for thesis adviser review.

Produces:
  - POLYCON_eval_results.xlsx   (metrics tables, definitions, human-rating template)
  - POLYCON_eval_appendix.docx  (side-by-side transcripts and summaries)
  - POLYCON_eval_artifacts.zip  (per-session pipeline output folders)
"""

from __future__ import annotations

import csv
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from eval.metrics import SessionMetrics, aggregate_metrics

HEADER_FILL = "1F4E79"
HEADER_FONT = "FFFFFF"
ALT_FILL = "F2F7FB"


@dataclass
class SessionEvalRecord:
    session_id: str
    title: str
    speaker_count: int
    transcription_enabled: bool
    reference_transcript: str
    reference_summary: str
    assemblyai_raw: str
    transcript_final: str
    notes_block: str
    summary_raw: str
    summary_sanitized: str
    metrics: SessionMetrics
    summary_upper_bound: str = ""
    artifact_dir: Path | None = None
    pipeline_meta: dict[str, Any] | None = None


def export_all(
    records: list[SessionEvalRecord],
    out_dir: Path,
    *,
    run_label: str = "",
) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "xlsx": out_dir / "POLYCON_eval_results.xlsx",
        "docx": out_dir / "POLYCON_eval_appendix.docx",
        "zip": out_dir / "POLYCON_eval_artifacts.zip",
        "csv": out_dir / "summary_report.csv",
    }
    export_csv(records, paths["csv"])
    export_excel(records, paths["xlsx"], run_label=run_label)
    export_word_appendix(records, paths["docx"], run_label=run_label)
    export_artifacts_zip(records, paths["zip"])
    return paths


def export_csv(records: list[SessionEvalRecord], path: Path) -> None:
    fieldnames = [
        "session_id",
        "title",
        "speaker_count",
        "wer_transcript_final_pct",
        "wer_transcript_final_fair_pct",
        "rougeL_pct",
        "semantic_similarity_pct",
        "rougeL_upper_pct",
        "semantic_similarity_upper_pct",
        "structured_mean_coverage_pct",
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in records:
            m = r.metrics.to_dict()
            writer.writerow(
                {
                    "session_id": r.session_id,
                    "title": r.title,
                    "speaker_count": r.speaker_count,
                    "wer_transcript_final_pct": m["wer_transcript_final_pct"],
                    "wer_transcript_final_fair_pct": m["wer_transcript_final_fair_pct"],
                    "rougeL_pct": m["rougeL_pct"],
                    "semantic_similarity_pct": m["semantic_similarity_pct"],
                    "rougeL_upper_pct": m["rougeL_upper_pct"],
                    "semantic_similarity_upper_pct": m["semantic_similarity_upper_pct"],
                    "structured_mean_coverage_pct": m["structured_mean_coverage_pct"],
                }
            )


def export_excel(records: list[SessionEvalRecord], path: Path, *, run_label: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color=HEADER_FONT)
    alt_fill = PatternFill("solid", fgColor=ALT_FILL)

    agg = aggregate_metrics([r.metrics for r in records])
    meta = records[0].pipeline_meta if records else {}

  # --- Executive Summary ---
    ws = wb.active
    ws.title = "Executive Summary"
    summary_rows = [
        ("POLYCON Consultation AI Evaluation", ""),
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Run label", run_label or "production-like"),
        ("Sessions evaluated", agg["session_count"]),
        ("", ""),
        ("TRANSCRIPTION (AssemblyAI + role labeling)", ""),
        ("Metric", "Value (%)"),
        ("Mean WER strict — final transcript", _fmt(agg["wer_transcript_final"]["mean"])),
        ("Mean WER fair/normalized — final transcript", _fmt(agg["wer_transcript_final_fair"]["mean"])),
        ("", ""),
        ("SUMMARIZATION — production pipeline", ""),
        ("Mean ROUGE-L", _fmt(agg["rougeL"]["mean"])),
        ("Mean semantic similarity", _fmt(agg["semantic_similarity"]["mean"])),
        ("Mean structured field coverage", _fmt(agg["structured_mean_coverage"]["mean"])),
        ("", ""),
        ("SUMMARIZATION — upper bound (reference transcript)", ""),
        ("Mean ROUGE-L upper bound", _fmt(agg["rougeL_upper"]["mean"])),
        ("Mean semantic similarity upper bound", _fmt(agg["semantic_similarity_upper"]["mean"])),
        ("", ""),
        ("MODEL CONFIG", ""),
        ("Gemini model", meta.get("gemini_model", "gemini-flash-lite-latest")),
        ("AssemblyAI speaker labels", "enabled"),
    ]
    for row_idx, (a, b) in enumerate(summary_rows, 1):
        ws.cell(row=row_idx, column=1, value=a)
        ws.cell(row=row_idx, column=2, value=b)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28

  # --- Transcription ---
    ws_t = wb.create_sheet("Transcription")
    t_headers = [
        "session_id",
        "title",
        "WER strict (%)",
        "WER fair/normalized (%)",
        "appendix_ref",
    ]
    _write_table(ws_t, t_headers, [
        [
            r.session_id,
            r.title,
            r.metrics.to_dict()["wer_transcript_final_pct"],
            r.metrics.to_dict()["wer_transcript_final_fair_pct"],
            f"Appendix — Session {r.session_id}",
        ]
        for r in records
    ], header_fill, header_font, alt_fill, border)

  # --- Summarization ---
    ws_s = wb.create_sheet("Summarization")
    s_headers = [
        "session_id",
        "title",
        "ROUGE-L prod (%)",
        "Semantic sim prod (%)",
        "ROUGE-L upper (%)",
        "Semantic sim upper (%)",
        "Struct coverage (%)",
    ]
    _write_table(ws_s, s_headers, [
        [
            r.session_id,
            r.title,
            r.metrics.to_dict()["rougeL_pct"],
            r.metrics.to_dict()["semantic_similarity_pct"],
            r.metrics.to_dict()["rougeL_upper_pct"],
            r.metrics.to_dict()["semantic_similarity_upper_pct"],
            r.metrics.to_dict()["structured_mean_coverage_pct"],
        ]
        for r in records
    ], header_fill, header_font, alt_fill, border)

  # --- Pipeline config ---
    ws_p = wb.create_sheet("Pipeline Config")
    p_headers = ["session_id", "title", "speaker_count", "transcription_enabled", "gemini_model", "artifact_folder"]
    _write_table(ws_p, p_headers, [
        [
            r.session_id,
            r.title,
            r.speaker_count,
            r.transcription_enabled,
            (r.pipeline_meta or {}).get("gemini_model", ""),
            f"session_{r.session_id}/",
        ]
        for r in records
    ], header_fill, header_font, alt_fill, border)

  # --- Definitions ---
    ws_d = wb.create_sheet("Definitions")
    defs = [
        ("WER strict", "Standard word error rate vs human reference transcript. Lower is better."),
        ("WER fair/normalized", "WER after normalizing fillers, contractions, and common tech-term variants (e.g. Socket.io vs Socket IO). Reported alongside strict WER."),
        ("ROUGE-L production", "Overlap for summary generated from AssemblyAI transcript + teacher notes."),
        ("Semantic similarity", "Embedding cosine similarity (all-MiniLM-L6-v2) — captures paraphrases ROUGE may miss."),
        ("Upper bound", "Summary generated from human reference transcript + notes — shows ceiling if transcription were perfect."),
        ("Structured coverage", "Share of significant concern/action/outcome terms reflected in the generated summary."),
        ("Human ratings", "Optional 1–5 Likert scores — fill in the Human Ratings sheet after faculty review."),
    ]
    for i, (term, desc) in enumerate(defs, 1):
        ws_d.cell(row=i, column=1, value=term).font = Font(bold=True)
        ws_d.cell(row=i, column=2, value=desc)
        ws_d.cell(row=i, column=2).alignment = Alignment(wrap_text=True)
    ws_d.column_dimensions["A"].width = 22
    ws_d.column_dimensions["B"].width = 72

  # --- Human ratings template ---
    ws_h = wb.create_sheet("Human Ratings")
    h_headers = [
        "session_id",
        "title",
        "rater_name",
        "factual_correctness_1_5",
        "coverage_1_5",
        "coherence_1_5",
        "alignment_1_5",
        "comments",
    ]
    _write_table(ws_h, h_headers, [[r.session_id, r.title, "", "", "", "", "", ""] for r in records],
                 header_fill, header_font, alt_fill, border)
    note_row = len(records) + 3
    ws_h.cell(row=note_row, column=1, value="Instructions: Ask 2–3 raters to score each summary. Leave blank until human review is done.")

    wb.save(path)


def _write_table(ws, headers, rows, header_fill, header_font, alt_fill, border):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for r_idx, row in enumerate(rows, 2):
        row_fill = alt_fill if r_idx % 2 == 0 else None
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            if row_fill is not None:
                cell.fill = row_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(36, len(headers[col - 1]) + 4))


def _fmt(value: float | None) -> str | float:
    if value is None:
        return "N/A"
    return value


def export_word_appendix(records: list[SessionEvalRecord], path: Path, *, run_label: str = "") -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    title = doc.add_heading("POLYCON Evaluation Appendix", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.add_run("Side-by-side reference (gold) vs system outputs for each consultation session.\n").bold = True
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    if run_label:
        doc.add_paragraph(f"Run: {run_label}")
    doc.add_paragraph(f"Sessions: {len(records)}")

    for r in records:
        doc.add_page_break()
        doc.add_heading(f"Session {r.session_id} — {r.title}", level=1)

        m = r.metrics.to_dict()
        stats = doc.add_paragraph()
        stats.add_run("Metrics: ").bold = True
        stats.add_run(
            f"WER strict {m['wer_transcript_final_pct']}% | "
            f"WER fair {m['wer_transcript_final_fair_pct']}% | "
            f"ROUGE-L {m['rougeL_pct']}% | "
            f"Semantic {m['semantic_similarity_pct']}% | "
            f"ROUGE-L upper {m['rougeL_upper_pct']}%"
        )

        doc.add_heading("Transcription", level=2)
        _add_comparison_table(doc, "Reference transcript (gold)", r.reference_transcript,
                              "System transcript (pipeline output)", r.transcript_final)

        doc.add_heading("AssemblyAI raw (before role labeling)", level=3)
        doc.add_paragraph(r.assemblyai_raw or "(empty)")

        doc.add_heading("Summarization", level=2)
        doc.add_paragraph("Notes block sent to Gemini:").runs[0].bold = True
        doc.add_paragraph(r.notes_block or "(empty)")
        _add_comparison_table(doc, "Reference summary (gold)", r.reference_summary,
                              "Generated summary (production pipeline)", r.summary_raw)
        if r.summary_upper_bound:
            doc.add_heading("Upper bound (reference transcript + notes)", level=3)
            doc.add_paragraph(r.summary_upper_bound)

    doc.save(path)


def _add_comparison_table(doc, left_title: str, left_text: str, right_title: str, right_text: str) -> None:
    from docx.shared import Pt

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = left_title
    hdr[1].text = right_title
    for cell in hdr:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    row = table.add_row().cells
    row[0].text = left_text or "(empty)"
    row[1].text = right_text or "(empty)"
    for cell in row:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(6)


def export_artifacts_zip(records: list[SessionEvalRecord], zip_path: Path) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for r in records:
            if r.artifact_dir and r.artifact_dir.is_dir():
                for file in r.artifact_dir.rglob("*"):
                    if file.is_file():
                        arcname = f"session_{r.session_id}/{file.relative_to(r.artifact_dir).as_posix()}"
                        zf.write(file, arcname)
        # Include machine-readable summary inside zip
        summary_csv = zip_path.parent / "summary_report.csv"
        if summary_csv.is_file():
            zf.write(summary_csv, "summary_report.csv")
